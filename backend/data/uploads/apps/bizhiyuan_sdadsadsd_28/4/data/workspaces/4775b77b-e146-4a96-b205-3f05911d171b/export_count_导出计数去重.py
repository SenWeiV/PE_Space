import argparse
import csv
import os
import sys
from collections import OrderedDict

try:
    import openpyxl
except Exception:
    openpyxl = None


def normalize_headers(raw_headers):
    headers = []
    used = {}
    for i, h in enumerate(raw_headers):
        name = str(h).strip() if h is not None else ""
        if name == "":
            name = f"Column{i + 1}"
        base = name
        if base in used:
            used[base] += 1
            name = f"{base}_{used[base]}"
        else:
            used[base] = 1
        headers.append(name)
    return headers


def read_csv(path):
    with open(path, newline="", encoding="utf-8-sig") as f:
        reader = csv.reader(f)
        rows = list(reader)
    if not rows:
        return [], []
    headers = normalize_headers(rows[0])
    data_rows = []
    for row in rows[1:]:
        row_dict = {}
        for i, header in enumerate(headers):
            row_dict[header] = row[i] if i < len(row) else ""
        data_rows.append(row_dict)
    return headers, data_rows


def read_xlsx(path, sheet_name=None):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if sheet_name:
        if sheet_name not in wb.sheetnames:
            raise ValueError(f"工作表不存在: {sheet_name}")
        ws = wb[sheet_name]
    else:
        ws = wb[wb.sheetnames[0]]
    rows_iter = ws.iter_rows(values_only=True)
    try:
        raw_headers = next(rows_iter)
    except StopIteration:
        return [], []
    headers = normalize_headers(raw_headers)
    data_rows = []
    for row in rows_iter:
        row_dict = {}
        for i, header in enumerate(headers):
            value = row[i] if i < len(row) else ""
            row_dict[header] = "" if value is None else value
        data_rows.append(row_dict)
    return headers, data_rows


def count_rows(rows, key_columns):
    counts = OrderedDict()
    for row in rows:
        key = tuple(row.get(col, "") for col in key_columns)
        if key not in counts:
            counts[key] = {"row": row, "count": 0}
        counts[key]["count"] += 1
    return counts


def unique_column_name(existing, base_name):
    if base_name not in existing:
        return base_name
    i = 2
    while f"{base_name}_{i}" in existing:
        i += 1
    return f"{base_name}_{i}"


def write_xlsx(path, headers, rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "counts"
    ws.append(headers)
    for row in rows:
        ws.append(row)
    wb.save(path)


def parse_columns(columns_arg, headers):
    if not columns_arg:
        return headers
    columns = [c.strip() for c in columns_arg.split(",") if c.strip()]
    missing = [c for c in columns if c not in headers]
    if missing:
        raise ValueError(f"列名不存在: {', '.join(missing)}")
    return columns


def ensure_openpyxl():
    if openpyxl is None:
        print("缺少依赖 openpyxl，请先安装：pip install openpyxl", file=sys.stderr)
        sys.exit(1)


def main():
    parser = argparse.ArgumentParser(description="导出计数并保留非重复项")
    parser.add_argument("input_path", help="输入 CSV 或 XLSX 文件路径")
    parser.add_argument("-o", "--output", help="输出 XLSX 文件路径")
    parser.add_argument("--sheet", help="指定 XLSX 工作表名称")
    parser.add_argument("--columns", help="用于计数的列名（逗号分隔），默认全部列")
    parser.add_argument("--output-all-columns", action="store_true", help="输出所有列")
    args = parser.parse_args()

    input_path = os.path.abspath(args.input_path)
    if not os.path.exists(input_path):
        print(f"输入文件不存在: {input_path}", file=sys.stderr)
        sys.exit(1)

    ext = os.path.splitext(input_path)[1].lower()
    ensure_openpyxl()
    if ext == ".csv":
        headers, rows = read_csv(input_path)
    elif ext == ".xlsx":
        headers, rows = read_xlsx(input_path, args.sheet)
    else:
        print("仅支持 CSV 或 XLSX 文件", file=sys.stderr)
        sys.exit(1)

    if not headers:
        print("未检测到表头，无法处理", file=sys.stderr)
        sys.exit(1)

    key_columns = parse_columns(args.columns, headers)
    counts = count_rows(rows, key_columns)

    output_columns = headers if args.output_all_columns else key_columns
    count_column = unique_column_name(output_columns, "count")
    output_headers = list(output_columns) + [count_column]

    output_rows = []
    for entry in counts.values():
        row_values = [entry["row"].get(col, "") for col in output_columns]
        row_values.append(entry["count"])
        output_rows.append(row_values)

    output_path = args.output
    if not output_path:
        base, _ = os.path.splitext(input_path)
        output_path = f"{base}_count.xlsx"
    if not output_path.lower().endswith(".xlsx"):
        output_path += ".xlsx"

    write_xlsx(output_path, output_headers, output_rows)
    print(f"已输出: {output_path}")


if __name__ == "__main__":
    main()
